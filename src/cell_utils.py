"""
单元格操作模块 - 负责单元格值、公式、样式的深拷贝

职责单一：只做单元格级别的读写操作，不涉及业务逻辑。
依赖：openpyxl（外部库）
"""

import re
import openpyxl
from openpyxl.worksheet.worksheet import Worksheet
from openpyxl.utils import get_column_letter


def copy_cell_style(src_cell: openpyxl.cell.Cell, target_cell: openpyxl.cell.Cell) -> None:
    """
    深拷贝单元格样式（字体、边框、填充、对齐、数字格式、保护）。
    不复制值和公式，仅复制视觉格式。
    """
    if src_cell.has_style:
        target_cell.font = src_cell.font.copy()
        target_cell.border = src_cell.border.copy()
        target_cell.fill = src_cell.fill.copy()
        target_cell.number_format = src_cell.number_format
        target_cell.protection = src_cell.protection.copy()
        target_cell.alignment = src_cell.alignment.copy()


def copy_cell_value(src_cell: openpyxl.cell.Cell, target_cell: openpyxl.cell.Cell) -> None:
    """
    复制单元格的值或公式。
    若值以 '=' 开头则视为公式，原样保留；否则复制普通值。
    """
    target_cell.value = src_cell.value


def copy_cell(src_cell: openpyxl.cell.Cell, target_cell: openpyxl.cell.Cell) -> None:
    """完整复制单元格：值/公式 + 样式 + 批注"""
    copy_cell_value(src_cell, target_cell)
    copy_cell_style(src_cell, target_cell)
    if src_cell.comment is not None:
        try:
            target_cell.comment = src_cell.comment.copy()
        except Exception:
            target_cell.comment = src_cell.comment


def copy_row(src_ws: Worksheet, src_row: int,
             target_ws: Worksheet, target_row: int,
             max_col: int, row_offset: int = 0,
             adjust_formulas: bool = False) -> None:
    """
    复制整行数据（含样式）从源Sheet到目标Sheet。
    """
    src_height = src_ws.row_dimensions[src_row].height
    if src_height is not None:
        target_ws.row_dimensions[target_row].height = src_height

    for col in range(1, max_col + 1):
        src_cell = src_ws.cell(row=src_row, column=col)
        tgt_cell = target_ws.cell(row=target_row, column=col)
        copy_cell_value(src_cell, tgt_cell)
        copy_cell_style(src_cell, tgt_cell)
        if src_cell.comment is not None:
            try:
                tgt_cell.comment = src_cell.comment.copy()
            except Exception:
                tgt_cell.comment = src_cell.comment

        if adjust_formulas:
            _translate_formula_indirect(tgt_cell, src_row)


def copy_column_widths(src_ws: Worksheet, target_ws: Worksheet, max_col: int) -> None:
    """
    将源Sheet的列宽复制到目标Sheet。
    仅在第一次写入表头时调用。
    """
    for col in range(1, max_col + 1):
        col_letter = get_column_letter(col)
        src_width = src_ws.column_dimensions[col_letter].width
        if src_width is not None:
            target_ws.column_dimensions[col_letter].width = src_width


def copy_merged_cells(src_ws: Worksheet, target_ws: Worksheet,
                      src_row_start: int, src_row_end: int,
                      target_row_start: int, row_offset: int) -> None:
    """
    复制源Sheet中指定行范围内的合并单元格信息到目标Sheet。
    仅复制完全在 [src_row_start, src_row_end] 范围内的合并区域。
    """
    from openpyxl.worksheet.cell_range import CellRange

    for merged_range in src_ws.merged_cells.ranges:
        mr_min = merged_range.min_row
        mr_max = merged_range.max_row
        # 仅复制完全在范围内的合并区域
        if mr_min >= src_row_start and mr_max <= src_row_end:
            new_min_row = mr_min + row_offset
            new_max_row = mr_max + row_offset
            new_range = CellRange(
                min_col=merged_range.min_col,
                max_col=merged_range.max_col,
                min_row=new_min_row,
                max_row=new_max_row,
            )
            target_ws.merge_cells(str(new_range))


_CELL_REF_RE = re.compile(r'^(\$?)([A-Z]{1,3})(\$?)(\d+)$')


def _translate_ref(part: str, src_row: int) -> str:
    """把单个 A1 引用改写成 INDIRECT+ROW() 的动态绑定。

    - 绝对行引用（如 $AA$1）保持不动；
    - 相对行引用（如 H7）按「源行 = src_row」计算偏移，绑定到当前所在行：
      H7 在源第7行 → INDIRECT("H"&ROW())；
      H5 在源第6行 → INDIRECT("H"&(ROW()-1))。
    非普通 A1 引用（含 ! 或 [ 的外部引用）原样返回。
    """
    m = _CELL_REF_RE.match(part)
    if not m:
        return part
    _col_abs, col_letters, row_abs, row_digits = m.groups()
    if row_abs:
        return part
    offset = int(row_digits) - src_row
    if offset == 0:
        row_expr = "ROW()"
    elif offset > 0:
        row_expr = f"(ROW()+{offset})"
    else:
        row_expr = f"(ROW()-{abs(offset)})"
    return f'INDIRECT("{col_letters}"&{row_expr})'


def _translate_ref_token(value: str, src_row: int) -> str:
    """处理一个引用操作数（单格如 H7，或范围如 H7:AL7）。"""
    if '!' in value or '[' in value:
        return value
    if ':' in value:
        left, right = value.split(':', 1)
        return _translate_ref(left, src_row) + ':' + _translate_ref(right, src_row)
    return _translate_ref(value, src_row)


def _translate_formula_indirect(cell: openpyxl.cell.Cell, src_row: int) -> None:
    """把单元格内所有「带具体行号」的相对引用改写为 INDIRECT+ROW() 动态绑定。

    这样公式无论复制到输出表的哪一行（或上方增删行），都会自动引用当前所在行的
    对应列（H~AL 等），不再受原文件行号影响；绝对行引用（$A$1 之类）保持原样。
    解析失败或无需改写时保留原始公式。
    """
    if cell.value is None or not isinstance(cell.value, str) or not cell.value.startswith('='):
        return
    original = cell.value
    try:
        from openpyxl.formula import Tokenizer
        tokens = Tokenizer(original).items
        parts = []
        changed = False
        for tok in tokens:
            if tok.type == 'OPERAND' and tok.subtype == 'RANGE':
                new_val = _translate_ref_token(tok.value, src_row)
                if new_val != tok.value:
                    changed = True
                parts.append(new_val)
            else:
                parts.append(tok.value)
        if changed:
            cell.value = '=' + ''.join(parts)
    except Exception:
        # 解析失败时保留原始公式，避免数据丢失
        cell.value = original


def is_effective_row(ws: Worksheet, row: int, max_col: int) -> bool:
    """判断一行是否包含可输出的真实数据：A列单独有值不算"""
    for col in range(2, max_col + 1):
        val = ws.cell(row=row, column=col).value
        if val is not None and str(val).strip() != '':
            return True
    return False


def should_copy_row(ws: Worksheet, row: int, max_col: int) -> bool:
    """判断一行是否应复制到输出：至少要有B列及之后的有效内容"""
    return is_effective_row(ws, row, max_col)


def is_row_empty(ws: Worksheet, row: int, max_col: int) -> bool:
    """判断指定行是否为空行（A列单独有值也视为空行）"""
    return not should_copy_row(ws, row, max_col)