"""检查源文件表尾结构 - 调试工具"""
import openpyxl
import os
import glob
import sys

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from sheet_analyzer import _detect_footer_start, find_real_last_row, find_real_max_col

raw_dir = r'd:\Code_Project\Python\xlsx_merge\data\raw_xlsx'
files = sorted(glob.glob(os.path.join(raw_dir, '*.xlsx')))

for fpath in files:
    fname = os.path.basename(fpath)
    if fname.startswith('~$'):
        continue
    wb = openpyxl.load_workbook(fpath)
    for sname in wb.sheetnames:
        ws = wb[sname]
        print(f'\n=== {fname} -> {sname} ===')
        
        # 确定真实边界
        real_max_row = find_real_last_row(ws, ws.max_column, ws.max_row)
        real_max_col = find_real_max_col(ws, real_max_row)
        print(f'  max_row={ws.max_row}, real_last_row={real_max_row}, real_max_col={real_max_col}')
        
        # 显示倒数15行内容
        start = max(1, real_max_row - 14)
        print('  倒数行内容:')
        for row in range(start, real_max_row + 1):
            vals = []
            for col in range(1, min(20, real_max_col + 1)):
                v = ws.cell(row=row, column=col).value
                if v is not None:
                    vals.append(f'{col}:{repr(str(v)[:25])}')
            if vals:
                print(f'    行{row}: {" | ".join(vals)}')
            else:
                print(f'    行{row}: (空行)')
        
        # 智能检测表尾
        header_end = 6
        footer_start, match_info = _detect_footer_start(ws, header_end, real_max_row, real_max_col)
        print(f'  检测结果: {match_info}')
        print(f'  表尾起始行: {footer_start}, 表尾结束行: {real_max_row}')
        
        # 尾部合并单元格
        print('  尾部合并单元格:')
        for mr in ws.merged_cells.ranges:
            if mr.min_row >= real_max_row - 10:
                print(f'    {mr}')