"""
Sheet结构分析模块 - 负责检测Sheet的真实数据边界与分区

职责：分析单个Sheet的结构，返回表头/表体/表尾的行范围。
不依赖合并逻辑，仅做结构识别。
"""

import re
from dataclasses import dataclass
from typing import Tuple, List
from openpyxl.worksheet.worksheet import Worksheet
from cell_utils import is_row_empty


@dataclass
class SheetRegion:
    """Sheet的行分区结果"""
    header_start: int
    header_end: int
    body_start: int
    body_end: int
    footer_start: int
    footer_end: int
    real_max_row: int
    real_max_col: int
    match_info: str = ""


SIGNATURE_KEYWORDS: Tuple[str, ...] = (
    "制表/日期",
    "部门审核/日期",
    "副总经理审核",
    "综合管理部审核",
    "物业总经理/日期",
)

LEGEND_KEYWORDS: Tuple[str, ...] = (
    "法定假日", "加班转调休", "休息日", "补休", "婚假",
    "正常班", "早班", "中班", "晚班",
    "法定假日加班", "出差", "年假", "工伤", "产假", "陪产假",
    "特批有薪假", "特批无薪假", "病假", "事假", "旷工", "丧假",
    "用餐",
)

FOOTER_SCAN_ROWS = 10


def _extract_month_token(text: str) -> str:
    m = re.search(r'(\d{4}年\d{1,2}月|\d{1,2}月)', text)
    return m.group(1) if m else ""


def _is_header_title_row(text: str, month_token: str) -> bool:
    return bool(text and "考勤表" in text and (not month_token or month_token in text))


def _is_calendar_row(text: str, month_token: str) -> bool:
    if not text:
        return False
    if month_token and month_token not in text:
        return False
    return ("WEEKDAY(" in text) or ("CHOOSE(" in text) or any(day in text for day in ["周日", "周一", "周二", "周三", "周四", "周五", "周六"])


def _get_row_text(ws: Worksheet, row: int, max_col: int) -> str:
    parts = []
    for col in range(1, max_col + 1):
        val = ws.cell(row=row, column=col).value
        if val is not None:
            parts.append(str(val).strip())
    return " ".join(parts)


def find_real_last_row(ws: Worksheet, max_col: int, reported_max_row: int) -> int:
    for row in range(reported_max_row, 0, -1):
        if not is_row_empty(ws, row, max_col):
            return row
    return 0


def find_real_max_col(ws: Worksheet, max_row: int) -> int:
    real_max = 0
    for row in range(1, max_row + 1):
        for col in range(ws.max_column, 0, -1):
            val = ws.cell(row=row, column=col).value
            if val is not None and str(val).strip() != '':
                if col > real_max:
                    real_max = col
                break
    return real_max if real_max > 0 else ws.max_column


def _detect_footer_start(ws: Worksheet, header_end: int, real_max_row: int, real_max_col: int) -> Tuple[int, str]:
    last_content_row = real_max_row
    while last_content_row > header_end and is_row_empty(ws, last_content_row, real_max_col):
        last_content_row -= 1
    if last_content_row <= header_end:
        return header_end + 1, "无数据行"

    scan_start = max(header_end + 1, last_content_row - FOOTER_SCAN_ROWS + 1)
    rows: List[Tuple[int, str]] = []
    for row in range(scan_start, last_content_row + 1):
        rows.append((row, _get_row_text(ws, row, real_max_col)))

    sig_row = None
    for row, text in rows:
        if any(kw in text for kw in SIGNATURE_KEYWORDS):
            sig_row = row
            break
    if sig_row is None:
        return last_content_row + 1, f"倒数{FOOTER_SCAN_ROWS}行内未找到签批行"

    legend_rows = [row for row, text in rows if any(kw in text for kw in LEGEND_KEYWORDS)]
    if not legend_rows:
        return last_content_row + 1, f"倒数{FOOTER_SCAN_ROWS}行内未找到图例行"

    legend_rows.sort()
    blocks = []
    start = prev = legend_rows[0]
    for row in legend_rows[1:]:
        if row == prev + 1:
            prev = row
        else:
            blocks.append((start, prev))
            start = prev = row
    blocks.append((start, prev))

    for start, end in reversed(blocks):
        if end >= last_content_row - 1 and (end - start + 1) >= 2:
            footer_start = min(sig_row, start)
            return footer_start, f"签批行{sig_row} + 图例行{footer_start}~{end}"

    return last_content_row + 1, f"签批行{sig_row}存在，但未找到连续图例块"


def analyze_sheet(ws: Worksheet, header_rows: int, footer_rows: int, trim_trailing: bool = True) -> SheetRegion:
    reported_max_row = ws.max_row
    initial_max_col = ws.max_column

    if trim_trailing:
        real_max_row = find_real_last_row(ws, initial_max_col, reported_max_row)
        real_max_col = find_real_max_col(ws, real_max_row)
    else:
        real_max_row = reported_max_row
        real_max_col = initial_max_col

    if real_max_row == 0:
        return SheetRegion(1, 0, 1, 0, 1, 0, 0, real_max_col)

    header_start = 1
    header_end = min(header_rows, real_max_row)

    title_row = None
    calendar_row = None
    month_token = ""
    header_scan_end = min(header_end + 6, real_max_row)
    for row in range(1, header_scan_end + 1):
        text = _get_row_text(ws, row, real_max_col)
        if not month_token:
            month_token = _extract_month_token(text)
        if title_row is None and _is_header_title_row(text, month_token):
            title_row = row
        if calendar_row is None and _is_calendar_row(text, month_token):
            calendar_row = row

    footer_start, match_info = _detect_footer_start(ws, header_end, real_max_row, real_max_col)
    footer_end = real_max_row

    body_start = header_end + 1
    body_end = footer_start - 1

    if title_row is not None and title_row < body_start:
        header_end = max(header_end, title_row)
        body_start = header_end + 1
    if calendar_row is not None and calendar_row < body_start:
        header_end = max(header_end, calendar_row)
        body_start = header_end + 1

    return SheetRegion(
        header_start=header_start,
        header_end=header_end,
        body_start=body_start,
        body_end=body_end,
        footer_start=footer_start,
        footer_end=footer_end,
        real_max_row=real_max_row,
        real_max_col=real_max_col,
        match_info=match_info,
    )


def describe_region(region: SheetRegion, file_name: str, sheet_name: str) -> str:
    lines = [
        f"  [{file_name} → {sheet_name}]",
        f"    真实数据行: {region.real_max_row}, 最大列: {region.real_max_col}",
        f"    表头: 第{region.header_start}~{region.header_end}行",
        f"    表体: 第{region.body_start}~{region.body_end}行",
        f"    表尾: 第{region.footer_start}~{region.footer_end}行",
    ]
    if region.match_info:
        lines.append(f"    匹配信息: {region.match_info}")
    return "\n".join(lines)
