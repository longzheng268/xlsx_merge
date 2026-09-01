"""检查源文件中的公式情况"""
import openpyxl, os

input_dir = os.path.join(os.path.dirname(os.path.dirname(__file__)), 'data', 'raw_xlsx')
files = [f for f in os.listdir(input_dir) if f.endswith('.xlsx') and not f.startswith('~$')]
files.sort()

for fname in files:
    fpath = os.path.join(input_dir, fname)
    wb = openpyxl.load_workbook(fpath, data_only=False)
    for sname in wb.sheetnames:
        ws = wb[sname]
        formulas = []
        for r in range(1, ws.max_row + 1):
            for c in range(1, ws.max_column + 1):
                cell = ws.cell(r, c)
                if isinstance(cell.value, str) and cell.value.startswith('='):
                    formulas.append((cell.coordinate, cell.value))
        print(f"\n{fname} -> {sname}: {len(formulas)} formulas")
        for coord, val in formulas[:20]:
            print(f"  {coord}: {val}")
        if len(formulas) > 20:
            print(f"  ... and {len(formulas)-20} more")
    wb.close()