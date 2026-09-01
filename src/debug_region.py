"""调试分区分析"""
import openpyxl
import sys
import os

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))

from cell_utils import is_row_empty

fpath = r'd:\Code_Project\Python\xlsx_merge\data\raw_xlsx\员工考勤表模板（7月份）.xlsx'
wb = openpyxl.load_workbook(fpath)
ws = wb[wb.sheetnames[1]]  # 非安防部门

header_rows = 6
max_col = ws.max_column

print(f'Sheet: {ws.title}, max_row={ws.max_row}, max_col={max_col}')

# find_real_last_row
real_max_row = ws.max_row
for row in range(ws.max_row, 0, -1):
    if not is_row_empty(ws, row, max_col):
        real_max_row = row
        break

print(f'\nreal_max_row={real_max_row}')

# 显示所有行
for row in range(1, real_max_row + 1):
    vals = []
    for col in range(1, min(20, max_col + 1)):
        v = ws.cell(row=row, column=col).value
        if v is not None:
            vals.append(f'{col}:{repr(str(v)[:25])}')
    empty = is_row_empty(ws, row, max_col)
    status = "(无内容)" if not vals else ""
    print(f'  行{row}: empty={empty} {status}{" | ".join(vals)}')

# 模拟智能表尾检测
print('\n--- 智能表尾检测过程 ---')
header_end = min(header_rows, real_max_row)
last_content_row = real_max_row
while last_content_row > header_end and is_row_empty(ws, last_content_row, max_col):
    last_content_row -= 1
print(f'last_content_row={last_content_row}')

footer_start = last_content_row
empty_streak = 0
for row in range(last_content_row, header_end, -1):
    if is_row_empty(ws, row, max_col):
        empty_streak += 1
        print(f'  行{row}: 空行, streak={empty_streak}')
        if empty_streak >= 2:
            footer_start = row + empty_streak
            print(f'  -> footer_start={footer_start}')
            break
    else:
        empty_streak = 0
        footer_start = row
        print(f'  行{row}: 有内容, streak=0, footer_start={row}')

print(f'\n最终: header_end={header_end}, footer_start={footer_start}')
print(f'body={header_end+1}~{footer_start-1}')