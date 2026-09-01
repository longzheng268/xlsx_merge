import openpyxl
import os

input_dir = os.path.join(os.path.dirname(os.path.dirname(__file__)), 'data', 'raw_xlsx')
files = [f for f in os.listdir(input_dir) if f.endswith('.xlsx') and not f.startswith('~$')]
files.sort()

for fname in files:
    fpath = os.path.join(input_dir, fname)
    print(f"\n===== {fname} =====")
    wb = openpyxl.load_workbook(fpath, data_only=False)
    print(f"  Sheets: {wb.sheetnames}")
    for sname in wb.sheetnames:
        ws = wb[sname]
        print(f"  --- Sheet: {sname} ---")
        print(f"  Max row: {ws.max_row}, Max col: {ws.max_column}")
        # Print first 5 rows
        for r in range(1, min(6, ws.max_row+1)):
            vals = [ws.cell(r, c).value for c in range(1, min(ws.max_column+1, 12))]
            print(f"  Row {r}: {vals}")
        # Print last 3 rows
        print(f"  ... (rows {6} to {max(6, ws.max_row-3)})")
        for r in range(max(6, ws.max_row-2), ws.max_row+1):
            vals = [ws.cell(r, c).value for c in range(1, min(ws.max_column+1, 12))]
            print(f"  Row {r}: {vals}")
    wb.close()