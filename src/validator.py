"""
数据校验模块 - 负责单元格级别的数据类型与一致性校验

职责：遍历数据行，按规则校验，生成精准的错误定位信息。
支持严格模式（遇错即停）和日志模式（收集全部错误）。
"""

from dataclasses import dataclass, field
from typing import Dict, Tuple, Type, List, Optional
from openpyxl.worksheet.worksheet import Worksheet
from openpyxl.utils import get_column_letter

from sheet_analyzer import SIGNATURE_KEYWORDS, LEGEND_KEYWORDS


@dataclass
class ValidationError:
    """单个校验错误"""
    file_name: str       # 文件名
    sheet_name: str      # Sheet名
    cell_coord: str      # 单元格坐标 (如 "C15")
    column: int          # 列号 (1-based)
    row: int             # 行号 (1-based)
    expected_types: str  # 期望类型描述
    actual_value: str    # 实际值
    reason: str          # 错误原因

    def __str__(self) -> str:
        return (
            f"[数据异常] 文件: '{self.file_name}' | Sheet: '{self.sheet_name}' | "
            f"位置: {self.cell_coord} | 原因: {self.reason}"
        )


@dataclass
class ValidationResult:
    """校验结果集合"""
    errors: List[ValidationError] = field(default_factory=list)
    total_rows_checked: int = 0
    total_cells_checked: int = 0
    semantic_warnings: List[str] = field(default_factory=list)

    @property
    def has_errors(self) -> bool:
        return len(self.errors) > 0

    @property
    def error_count(self) -> int:
        return len(self.errors)

    def summary(self) -> str:
        """生成校验结果摘要"""
        lines = [
            f"校验完成: 共检查 {self.total_rows_checked} 行, {self.total_cells_checked} 个单元格",
            f"发现 {self.error_count} 个异常",
        ]
        for warn in self.semantic_warnings:
            lines.append(warn)
        for err in self.errors:
            lines.append(str(err))
        return "\n".join(lines)


def _row_text(ws: Worksheet, row: int, max_col: int) -> str:
    parts = []
    for col in range(1, max_col + 1):
        val = ws.cell(row=row, column=col).value
        if val is not None and str(val).strip() != "":
            parts.append(str(val).strip())
    return " ".join(parts)


def _is_header_like(text: str) -> bool:
    return bool(text and ("考勤表" in text or "日期" in text or "周" in text or any(day in text for day in ["周日", "周一", "周二", "周三", "周四", "周五", "周六"])))


def _is_footer_like(text: str) -> bool:
    return bool(text and (any(kw in text for kw in SIGNATURE_KEYWORDS) or any(kw in text for kw in LEGEND_KEYWORDS)))


def _is_serial_only(text: str, ws: Worksheet, row: int, max_col: int) -> bool:
    if not text:
        return False
    if len(text) > 4:
        return False
    if any(ch.isdigit() for ch in text):
        for col in range(2, max_col + 1):
            val = ws.cell(row=row, column=col).value
            if val is not None and str(val).strip() != "":
                return False
        return True
    return False


def validate_cell(value, expected_types: Tuple[Type, ...],
                  file_name: str, sheet_name: str,
                  row: int, col: int) -> Optional[ValidationError]:
    """
    校验单个单元格的值是否符合期望类型。
    """
    if value is None:
        return None
    if isinstance(value, str) and value.startswith('='):
        return None

    if not isinstance(value, expected_types):
        cell_coord = f"{get_column_letter(col)}{row}"
        type_names = "或".join(t.__name__ for t in expected_types)
        actual_type = type(value).__name__
        reason = f"期待类型: {type_names}, 实际类型: {actual_type}, 实际值: '{value}'"
        return ValidationError(
            file_name=file_name,
            sheet_name=sheet_name,
            cell_coord=cell_coord,
            column=col,
            row=row,
            expected_types=type_names,
            actual_value=str(value),
            reason=reason,
        )
    return None


def validate_sheet(ws: Worksheet, file_name: str, sheet_name: str,
                   body_start: int, body_end: int,
                   validation_rules: Dict[int, Tuple[Type, ...]],
                   strict_mode: bool = True) -> ValidationResult:
    """校验Sheet中数据体的所有单元格，并做 body 语义扫描。"""
    result = ValidationResult()

    if body_start > body_end:
        return result

    max_col = ws.max_column

    for row in range(body_start, body_end + 1):
        result.total_rows_checked += 1
        text = _row_text(ws, row, max_col)
        if _is_header_like(text):
            msg = f"[语义异常] 文件: '{file_name}' | Sheet: '{sheet_name}' | 行 {row} 混入表头内容: {text[:120]}"
            result.semantic_warnings.append(msg)
        if _is_footer_like(text):
            msg = f"[语义异常] 文件: '{file_name}' | Sheet: '{sheet_name}' | 行 {row} 混入表尾内容: {text[:120]}"
            result.semantic_warnings.append(msg)
        if _is_serial_only(text, ws, row, max_col):
            msg = f"[语义异常] 文件: '{file_name}' | Sheet: '{sheet_name}' | 行 {row} 疑似只有序号无实质内容: {text[:120]}"
            result.semantic_warnings.append(msg)

        for col, expected_types in validation_rules.items():
            result.total_cells_checked += 1
            cell_value = ws.cell(row=row, column=col).value
            error = validate_cell(cell_value, expected_types, file_name, sheet_name, row, col)
            if error is not None:
                result.errors.append(error)
                if strict_mode:
                    raise ValueError(str(error))

    return result


def validate_all_sheets(validation_results: List[ValidationResult]) -> str:
    """汇总所有Sheet的校验结果，生成报告。"""
    total_errors = sum(r.error_count for r in validation_results)
    total_rows = sum(r.total_rows_checked for r in validation_results)
    total_cells = sum(r.total_cells_checked for r in validation_results)
    total_warnings = sum(len(r.semantic_warnings) for r in validation_results)

    lines = [
        "=" * 60,
        "数据校验报告",
        "=" * 60,
        f"总检查行数: {total_rows}",
        f"总检查单元格数: {total_cells}",
        f"总异常数: {total_errors}",
        f"总语义告警数: {total_warnings}",
        "-" * 60,
    ]

    for r in validation_results:
        for warn in r.semantic_warnings:
            lines.append(warn)
        if r.has_errors:
            for err in r.errors:
                lines.append(str(err))

    if total_errors == 0 and total_warnings == 0:
        lines.append("✓ 所有数据校验通过，无异常。")

    lines.append("=" * 60)
    return "\n".join(lines)
