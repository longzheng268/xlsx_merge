import openpyxl

path = r'd:\Code_Project\Python\xlsx_merge\data\out_xlsx\merged_result.xlsx'
wb = openpyxl.load_workbook(path, data_only=False)
ws = wb[wb.sheetnames[0]]

print('sheet=', ws.title)
print('max_row=', ws.max_row, 'max_col=', ws.max_column)

for row in range(1, ws.max_row + 1):
    vals = []
    for col in range(1, min(20, ws.max_column) + 1):
        v = ws.cell(row=row, column=col).value
        if v is not None and str(v).strip() != '':
            vals.append(str(v).strip())
    text = ' | '.join(vals)
    if any(k in text for k in ['制表', '部门审核', '副总经理审核', '综合管理部审核', '物业总经理', '法定假日', '加班转调休', '正常班', '早班', '中班', '晚班']):
        print(f'ROW {row}: {text[:300]}')
